# SPDX-License-Identifier: AGPL-3.0-or-later
# Commercial license available
# © Concepts 1996–2026 Miroslav Šotek. All rights reserved.
# © Code 2020–2026 Miroslav Šotek. All rights reserved.
# ORCID: 0009-0009-3560-0851
# Contact: www.anulum.li | protoscience@anulum.li
# Atlas Reactorum — Swiss industrial producer conformance

"""Build all accepted Swiss facilities from original source cells and read-only snapshots."""

from __future__ import annotations

import io
import json
import shutil
import ssl
import subprocess
import threading
import time
import zipfile
from collections.abc import Iterator
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from types import ModuleType
from urllib.error import HTTPError, URLError

import openpyxl
import pytest

from ._catalogue_inputs import ROOT, read_table, run_cli, write_table
from .conftest import load_module

DIRECTORY = ROOT / "05_global_reactor_map/imports/industrial_facilities/expansion_round4"
SCRIPT = DIRECTORY / "build_dataset.py"
SNAPSHOT = DIRECTORY / "selected_source_snapshot.tsv"
DATA = DIRECTORY / "industrial_facilities_round4.tsv"
SOURCE = ROOT / "tests/data/industrial_round4/swiss_prtr.xlsx"


@pytest.fixture
def producer() -> ModuleType:
    """Load the Swiss workbook producer through its repository module path.

    Returns
    -------
    ModuleType
        Producer exposing the native selection, refresh and build entry points.
    """
    return load_module(str(SCRIPT.relative_to(ROOT)), "atlas_industrial_round4_build")


def raw_records() -> tuple[list[str], list[list[object]]]:
    """Read every cell from the accepted workbook's first sheet.

    Returns
    -------
    tuple[list[str], list[list[object]]]
        Column names and complete source rows, including repeated facilities.

    Notes
    -----
    Read cached cell values in read-only mode and close the workbook on exit.
    """
    workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    try:
        values = workbook.worksheets[0].iter_rows(values_only=True)
        fields = [str(v) for v in next(values)]
        return fields, [list(row) for row in values]
    finally:
        workbook.close()


def raw_payload(fields: list[str], rows: list[list[object]]) -> bytes:
    """Serialize the supplied complete source cells into an owned workbook copy.

    Parameters
    ----------
    fields : list[str]
        Source column names in their workbook order.
    rows : list[list[object]]
        Complete source rows with the case-specific cell mutations.

    Returns
    -------
    bytes
        XLSX content with one sheet; the original workbook is never written.
    """
    workbook = openpyxl.Workbook()
    worksheet = workbook.worksheets[0]
    worksheet.append(fields)
    for row in rows:
        worksheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def test_actual_raw_selection_equals_complete_snapshot(producer: ModuleType) -> None:
    """Select all 93 accepted facilities from 469 raw rows with exact snapshot equality."""
    selected, count = producer.select_source(SOURCE.read_bytes())
    assert count == 469 and selected == producer.read_snapshot(SNAPSHOT)
    assert len(selected) == 93
    assert producer.clean(None) == "" and producer.clean(" observed ") == "observed"


def test_native_full_rebuild_preserves_historical_bytes(
    tmp_path: Path, producer: ModuleType
) -> None:
    """Rebuild exact historical bytes through normal, optimized and API entry points."""
    before = SNAPSHOT.read_bytes(), DATA.read_bytes()
    for optimize in (False, True):
        output = tmp_path / f"built-{optimize}.tsv"
        result = run_cli(SCRIPT, "--output", str(output), optimize=optimize)
        assert result.returncode == 0, result.stdout + result.stderr
        assert output.read_bytes() == DATA.read_bytes()
    assert producer.build(SNAPSHOT, tmp_path / "api.tsv") == 93
    assert (tmp_path / "api.tsv").read_bytes() == DATA.read_bytes()
    assert before == (SNAPSHOT.read_bytes(), DATA.read_bytes())


def test_actual_raw_refresh_is_isolated(tmp_path: Path, producer: ModuleType) -> None:
    """Refresh owned copies through API and both CLI modes without changing accepted data."""
    snapshot, manifest = tmp_path / "snapshot.tsv", tmp_path / "manifest.tsv"
    assert producer.refresh_snapshot(
        SOURCE.read_bytes(), snapshot=snapshot, manifest=manifest, retrieved="2026-09-28"
    ) == (469, 93)
    assert snapshot.read_bytes() == SNAPSHOT.read_bytes()
    assert read_table(manifest)[1][0]["selected_facilities"] == "93"
    for optimize in (False, True):
        result = run_cli(
            SCRIPT,
            "--refresh",
            "--raw-source",
            str(SOURCE),
            "--snapshot",
            str(tmp_path / "cli-snapshot.tsv"),
            "--manifest",
            str(tmp_path / "cli-manifest.tsv"),
            "--output",
            str(tmp_path / "cli.tsv"),
            "--date",
            "2026-09-28",
            optimize=optimize,
        )
        assert result.returncode == 0, result.stdout + result.stderr
        assert (tmp_path / "cli.tsv").read_bytes() == DATA.read_bytes()


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("source_record_id", ""),
        ("source_record_id", "invalid"),
        ("industrial_sector", ""),
        ("source_north_lv95", ""),
        ("source_north_lv95", "unknown"),
        ("source_north_lv95", "999999"),
        ("source_east_lv95", "NaN"),
        ("facility_name", ""),
        ("operator", ""),
        ("source", "unknown"),
        ("reporting_year", "2023"),
        ("country", "unknown"),
        ("nace_code", "99"),
        ("lat", "unknown"),
        ("lat", "NaN"),
        ("lon", "181"),
    ],
)
def test_full_snapshot_mutations_refuse(tmp_path: Path, field: str, value: str) -> None:
    """Reject each invalid source observation without a traceback or output table."""
    fields, rows = read_table(SNAPSHOT)
    rows[0][field] = value
    path = tmp_path / "changed.tsv"
    write_table(path, fields, rows)
    result = run_cli(
        SCRIPT, "--snapshot", str(path), "--output", str(tmp_path / "output.tsv"), optimize=True
    )
    assert result.returncode == 1 and "FAIL" in result.stdout
    assert "Traceback" not in result.stderr and not (tmp_path / "output.tsv").exists()


@pytest.mark.parametrize(
    "corruption", ["empty", "header", "extra", "short", "quote", "utf8", "duplicate", "missing"]
)
def test_snapshot_integrity_refuses(tmp_path: Path, corruption: str) -> None:
    """Reject malformed, duplicate or missing snapshot input through the native CLI."""
    fields, rows = read_table(SNAPSHOT)
    path = tmp_path / "changed.tsv"
    if corruption == "empty":
        rows.clear()
    elif corruption == "header":
        fields = fields[::-1]
    elif corruption == "duplicate":
        rows.append(rows[0].copy())
    write_table(path, fields, rows)
    if corruption == "extra":
        path.write_text(path.read_text() + "\t".join(rows[0].values()) + "\textra\n")
    elif corruption == "short":
        path.write_text("\t".join(fields) + "\nonly-one-cell\n")
    elif corruption == "quote":
        path.write_text("\t".join(fields) + '\n"unterminated')
    elif corruption == "utf8":
        path.write_bytes(b"\xff")
    elif corruption == "missing":
        path.unlink()
    result = run_cli(SCRIPT, "--snapshot", str(path), "--output", str(tmp_path / "output.tsv"))
    assert result.returncode == 1 and "Traceback" not in result.stderr


def test_original_optional_activity_and_public_schema(producer: ModuleType) -> None:
    """Keep optional activities absent and reject empty or non-string snapshot schemas."""
    rows = producer.read_snapshot(SNAPSHOT)
    rows[0]["ordinance_level_2"] = ""
    rows[0]["ordinance_level_3"] = ""
    output = producer.rows_from_snapshot(rows)
    assert len(output) == 93 and "PRTRO level" not in output[0]["process_or_activity"]
    with pytest.raises(ValueError, match="no records"):
        producer.rows_from_snapshot([])
    changed = json.loads(json.dumps(rows))
    changed[0]["facility_name"] = None
    with pytest.raises(ValueError, match="schema"):
        producer.rows_from_snapshot(changed)
    rows[0].pop("country")
    with pytest.raises(ValueError, match="schema"):
        producer.rows_from_snapshot(rows)
    assert producer.clean(2016.0) == "2016" and producer.clean(2016.5) == "2016.5"


@pytest.mark.parametrize(
    "corruption",
    [
        "header",
        "duplicate_header",
        "empty",
        "inconsistent",
        "empty_id",
        "bad_id",
        "no_selection",
        "north",
        "east",
        "missing_point",
        "bad_number",
    ],
)
def test_actual_raw_boundaries_refuse(producer: ModuleType, corruption: str) -> None:
    """Reject invalid workbook identity, selection, coordinates and conflicting repeats."""
    fields, rows = raw_records()
    index = {field: i for i, field in enumerate(fields)}
    selected = rows[3]
    if corruption == "header":
        fields[0] = "unexpected"
    elif corruption == "duplicate_header":
        fields[0] = fields[1]
    elif corruption == "empty":
        rows.clear()
    elif corruption == "inconsistent":
        changed = selected.copy()
        changed[index["Owner"]] = "changed operator"
        rows.append(changed)
    elif corruption in {"empty_id", "bad_id"}:
        selected[index["Facility ID"]] = None if corruption == "empty_id" else "invalid"
    elif corruption == "no_selection":
        for row in rows:
            row[index["NACE code"]] = 99
    elif corruption == "north":
        selected[index["North coordinate (CH1903+)"]] = 999999
    elif corruption == "east":
        selected[index["East coordinate (CH1903+)"]] = float("inf")
    elif corruption == "missing_point":
        selected[index["East coordinate (CH1903+)"]] = None
    else:
        selected[index["East coordinate (CH1903+)"]] = "unknown"
    with pytest.raises(ValueError):
        producer.select_source(raw_payload(fields, rows))


def test_actual_archive_limits_and_corruption(producer: ModuleType) -> None:
    """Reject expanded workbook content above 64 MiB and truncated ZIP input."""
    stream = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(SOURCE.read_bytes())) as original:
        with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for item in original.infolist():
                archive.writestr(item, original.read(item.filename))
            archive.writestr("oversize-control.bin", b"0" * (64 * 1024 * 1024 + 1))
    with pytest.raises(ValueError, match="expanded content"):
        producer.select_source(stream.getvalue())
    with pytest.raises(zipfile.BadZipFile):
        producer.select_source(SOURCE.read_bytes()[:100])


@pytest.mark.parametrize("date", ["20260928", "2026-02-30"])
def test_exact_retrieval_dates_required(tmp_path: Path, producer: ModuleType, date: str) -> None:
    """Reject compact or impossible calendar dates in both public conversion paths."""
    with pytest.raises(ValueError):
        producer.rows_from_snapshot(producer.read_snapshot(SNAPSHOT), retrieved=date)
    with pytest.raises(ValueError):
        producer.refresh_snapshot(
            SOURCE.read_bytes(),
            snapshot=tmp_path / "snapshot.tsv",
            manifest=tmp_path / "manifest.tsv",
            retrieved=date,
        )


def test_historical_protection_and_atomic_failures(tmp_path: Path, producer: ModuleType) -> None:
    """Protect accepted inputs and preserve prior output when atomic serialization fails."""
    before = DATA.read_bytes(), SNAPSHOT.read_bytes()
    for target in [DATA, SNAPSHOT, tmp_path / "link.tsv"]:
        if target.name == "link.tsv":
            target.symlink_to(DATA)
        result = run_cli(SCRIPT, "--output", str(target))
        assert result.returncode == 1 and "separate" in result.stdout
    occupied = tmp_path / "occupied"
    occupied.write_text("preserve")
    assert producer.main(["--output", str(occupied / "output.tsv")]) == 1
    directory = tmp_path / "directory"
    directory.mkdir()
    assert producer.main(["--output", str(directory)]) == 1
    rows = producer.rows_from_snapshot(producer.read_snapshot(SNAPSHOT))
    rows[0]["facility_name"] = "\ud800"
    output = tmp_path / "atomic.tsv"
    output.write_text("preserve")
    with pytest.raises(UnicodeError):
        producer.write_table(output, producer.FIELDS, rows)
    assert output.read_text() == "preserve" and not list(tmp_path.glob("*.tmp"))
    assert before == (DATA.read_bytes(), SNAPSHOT.read_bytes())


def test_cli_refresh_refuses_bad_configuration(tmp_path: Path, producer: ModuleType) -> None:
    """Reject incomplete refresh paths, aliases and absent raw input through native entry points."""
    assert (
        producer.main(
            ["--output", str(tmp_path / "out.tsv"), "--manifest", str(tmp_path / "m.tsv")]
        )
        == 1
    )
    assert producer.main(["--refresh", "--output", str(tmp_path / "out.tsv")]) == 1
    assert (
        producer.main(
            [
                "--refresh",
                "--snapshot",
                str(tmp_path / "s.tsv"),
                "--manifest",
                str(tmp_path / "s.tsv"),
                "--output",
                str(tmp_path / "out.tsv"),
            ]
        )
        == 1
    )
    assert (
        producer.main(
            [
                "--refresh",
                "--snapshot",
                str(tmp_path / "s.tsv"),
                "--manifest",
                str(tmp_path / "m.tsv"),
                "--output",
                str(tmp_path / "out.tsv"),
                "--raw-source",
                str(tmp_path / "missing.csv"),
            ]
        )
        == 1
    )
    assert run_cli(SCRIPT).returncode == 2


class SourceServer(ThreadingHTTPServer):
    """Serve the complete genuine Swiss selection over a real trusted TLS socket."""

    body: bytes


class SourceHandler(BaseHTTPRequestHandler):
    """Provide real HTTP response, redirect and timeout controls."""

    def do_GET(self) -> None:
        """Serve source bytes, redirects, status failures and a delayed response over real TLS."""
        server = self.server
        assert isinstance(server, SourceServer)
        if self.path in {"/redirect", "/downgrade", "/credentials", "/loop"}:
            self.send_response(302)
            target = {
                "/redirect": "/csv",
                "/downgrade": "http://127.0.0.1/source",
                "/credentials": "https://user:secret@localhost/source",
                "/loop": "/loop",
            }[self.path]
            self.send_header("Location", target)
            self.end_headers()
            return
        self.send_response(
            201 if self.path == "/created" else 404 if self.path == "/missing" else 200
        )
        self.send_header("Content-Length", str(len(server.body)))
        self.end_headers()
        if self.path == "/stall":
            time.sleep(0.2)
        try:
            self.wfile.write(server.body)
        except (BrokenPipeError, ConnectionResetError, ssl.SSLError):
            pass

    def log_message(self, fmt: str, *args: object) -> None:
        """Keep native HTTP diagnostics out of successful test output."""


@pytest.fixture
def tls_source(tmp_path: Path) -> Iterator[tuple[str, Path]]:
    """Serve the original workbook over a temporary, locally trusted TLS listener.

    Parameters
    ----------
    tmp_path : Path
        Owned directory for the one-day certificate and private key.

    Yields
    ------
    tuple[str, Path]
        HTTPS origin and certificate path for the native downloader.

    Raises
    ------
    RuntimeError
        If the required OpenSSL executable is unavailable.

    Notes
    -----
    Stop and close the listener and join its worker after the test.
    """
    ca, key = tmp_path / "ca.pem", tmp_path / "key.pem"
    openssl = shutil.which("openssl")
    if openssl is None:
        raise RuntimeError("Native TLS conformance requires OpenSSL")
    subprocess.run(
        [
            openssl,
            "req",
            "-x509",
            "-newkey",
            "rsa:2048",
            "-nodes",
            "-keyout",
            str(key),
            "-out",
            str(ca),
            "-days",
            "1",
            "-subj",
            "/CN=localhost",
            "-addext",
            "subjectAltName=DNS:localhost,IP:127.0.0.1",
        ],
        check=True,
        capture_output=True,
        timeout=30,
    )
    server = SourceServer(("127.0.0.1", 0), SourceHandler)
    server.body = SOURCE.read_bytes()
    context = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
    context.load_cert_chain(ca, key)
    server.socket = context.wrap_socket(server.socket, server_side=True)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"https://127.0.0.1:{server.server_port}", ca
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_real_tls_and_complete_native_refresh(
    tmp_path: Path, producer: ModuleType, tls_source: tuple[str, Path]
) -> None:
    """Refresh all 93 facilities over trusted TLS and protect the certificate from output aliasing."""
    url, ca = tls_source
    assert producer.download(url + "/redirect", ca_file=ca) == SOURCE.read_bytes()
    result = run_cli(
        SCRIPT,
        "--refresh",
        "--snapshot",
        str(tmp_path / "snapshot.tsv"),
        "--manifest",
        str(tmp_path / "manifest.tsv"),
        "--output",
        str(tmp_path / "output.tsv"),
        "--source-url",
        url + "/csv",
        "--ca-file",
        str(ca),
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert len(read_table(tmp_path / "output.tsv")[1]) == 93
    assert (tmp_path / "snapshot.tsv").read_bytes() == SNAPSHOT.read_bytes()
    before = ca.read_bytes()
    assert (
        producer.main(
            [
                "--refresh",
                "--snapshot",
                str(tmp_path / "guard-snapshot.tsv"),
                "--manifest",
                str(tmp_path / "guard-manifest.tsv"),
                "--output",
                str(ca),
                "--ca-file",
                str(ca),
            ]
        )
        == 1
    )
    assert ca.read_bytes() == before


@pytest.mark.parametrize(
    "path", ["/downgrade", "/credentials", "/loop", "/missing", "/created", "/stall"]
)
def test_native_transport_refusals(
    producer: ModuleType, tls_source: tuple[str, Path], path: str
) -> None:
    """Reject unsafe redirects, non-200 responses and the actual delayed-response timeout."""
    url, ca = tls_source
    with pytest.raises((ValueError, HTTPError, URLError, TimeoutError)):
        producer.download(url + path, ca_file=ca, timeout=0.03 if path == "/stall" else 5)


def test_real_certificate_trust_and_limits(
    producer: ModuleType, tls_source: tuple[str, Path]
) -> None:
    """Reject untrusted TLS, oversized responses, invalid limits and unsafe source URLs."""
    url, ca = tls_source
    with pytest.raises(URLError):
        producer.download(url + "/csv")
    with pytest.raises(ValueError, match="byte limit"):
        producer.download(url + "/csv", ca_file=ca, max_bytes=10)
    for timeout, limit in [(0, 100), (float("inf"), 100), (1, 0)]:
        with pytest.raises(ValueError):
            producer.download(url + "/csv", timeout=timeout, max_bytes=limit)
    for invalid in [
        "http://localhost/source",
        "https:///source",
        "https://u:p@localhost/source",
        "https://localhost:0/source",
        "https://localhost:99999/source",
        "https://localhost/white space",
    ]:
        with pytest.raises(ValueError):
            producer.download(invalid)
