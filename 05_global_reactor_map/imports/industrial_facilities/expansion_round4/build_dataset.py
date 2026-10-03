#!/usr/bin/env python3
# SPDX-License-Identifier: AGPL-3.0-or-later
# Commercial license available
# © Concepts 1996–2026 Miroslav Šotek. All rights reserved.
# © Code 2020–2026 Miroslav Šotek. All rights reserved.
# ORCID: 0009-0009-3560-0851
# Contact: www.anulum.li | protoscience@anulum.li
# Atlas Reactorum — 05_global_reactor_map/imports/industrial_facilities/expansion_round4/build_dataset.py
"""Build the SwissPRTR industrial-facility discovery expansion.

Historical builds read the accepted compact snapshot and write a separate
explicit output. Refreshes preserve the documented SwissPRTR NACE selection
and coordinate transformation, with separate dated snapshot and manifest.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import math
import os
import ssl
import tempfile
import urllib.parse
import urllib.request
import zipfile
from datetime import date
from http.client import HTTPMessage
from pathlib import Path
from typing import IO

ROOT = Path(__file__).resolve().parent
SNAPSHOT = ROOT / "selected_source_snapshot.tsv"
MANIFEST = ROOT / "source_snapshot_manifest.tsv"
OUTPUT = ROOT / "industrial_facilities_round4.tsv"
RETRIEVED = "2026-09-28"
REPORTING_YEAR = "2024"
USER_AGENT = "reactor-atlas-industrial-discovery/4.0"

RAW_URL = "https://www.bafu.admin.ch/dam/en/sd-web/CXgeZKcf8ez8/swissprtr-daten-2007-2024.xlsx"
LANDING_URL = "https://www.bafu.admin.ch/en/swissprtr-en"
NACE_PREFIXES = ("10", "11", "17", "19", "20", "21", "22")

FIELDS = [
    "stable_id",
    "facility_name",
    "country",
    "lat",
    "lon",
    "precision",
    "sector",
    "process_or_activity",
    "reactor_type_if_explicit",
    "status",
    "operator",
    "capacity",
    "pollutant_or_product_context",
    "source_url",
    "source_role",
    "retrieved",
    "license",
    "verification_notes",
]
SNAPSHOT_FIELDS = [
    "source",
    "source_record_id",
    "reporting_year",
    "facility_name",
    "operator",
    "country",
    "lat",
    "lon",
    "source_north_lv95",
    "source_east_lv95",
    "nace_code",
    "industrial_sector",
    "ordinance_level_2",
    "ordinance_level_3",
]


MANIFEST_FIELDS = [
    "source",
    "url",
    "retrieved",
    "reporting_year",
    "bytes",
    "sha256",
    "raw_data_rows",
    "selected_facilities",
]
RAW_FIELDS = [
    "Year",
    "Source type",
    "Facility",
    "Owner",
    "Facility ID",
    "North coordinate (CH1903+)",
    "East coordinate (CH1903+)",
    "NACE code",
    "Industrial sector",
    "Ordinance level 2",
    "Ordinance level 3",
]


def check_https(url: str) -> None:
    """Require an anonymous HTTPS URL for every request and redirect."""
    parsed = urllib.parse.urlsplit(url)
    if (
        parsed.scheme != "https"
        or not parsed.hostname
        or parsed.username is not None
        or parsed.port == 0
        or any(c.isspace() for c in url)
    ):
        raise ValueError("source URL must be valid anonymous HTTPS")


class SourceRedirect(urllib.request.HTTPRedirectHandler):
    """Refuse non-HTTPS or credential-bearing redirect destinations."""

    max_redirections = 5

    def redirect_request(
        self,
        req: urllib.request.Request,
        fp: IO[bytes],
        code: int,
        msg: str,
        headers: HTTPMessage,
        newurl: str,
    ) -> urllib.request.Request | None:
        """Validate the real server's destination before urllib follows it."""
        check_https(newurl)
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def download(
    url: str,
    *,
    timeout: float = 60,
    max_bytes: int = 32 * 1024 * 1024,
    ca_file: Path | None = None,
) -> bytes:
    """Read a bounded HTTPS resource with certificate verification.

    The timeout bounds socket operations, not the entire download duration.
    An optional CA file supports a trusted institutional mirror.
    """
    check_https(url)
    if not math.isfinite(timeout) or timeout <= 0 or max_bytes <= 0:
        raise ValueError("source limits must be positive and finite")
    context = ssl.create_default_context(cafile=str(ca_file) if ca_file else None)
    opener = urllib.request.build_opener(
        urllib.request.HTTPSHandler(context=context),
        SourceRedirect(),
    )
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with opener.open(request, timeout=timeout) as response:
        if response.status != 200:
            raise ValueError(f"source returned HTTP {response.status}")
        body: bytes = response.read(max_bytes + 1)
    if len(body) > max_bytes:
        raise ValueError("source exceeds the download byte limit")
    return body


def clean(value: object) -> str:
    """Coerce a source value to trimmed text."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def lv95_to_wgs84(east: float, north: float) -> tuple[float, float]:
    """Apply swisstopo's December 2016 direct approximate formula."""
    y = (east - 2_600_000.0) / 1_000_000.0
    x = (north - 1_200_000.0) / 1_000_000.0
    longitude = (
        (2.6779094 + 4.728982 * y + 0.791484 * y * x + 0.1306 * y * x * x - 0.0436 * y * y * y)
        * 100.0
        / 36.0
    )
    latitude = (
        (
            16.9023892
            + 3.238272 * x
            - 0.270978 * y * y
            - 0.002528 * x * x
            - 0.0447 * y * y * x
            - 0.0140 * x * x * x
        )
        * 100.0
        / 36.0
    )
    return latitude, longitude


def read_snapshot(path: Path) -> list[dict[str, str]]:
    """Read the exact field-minimised snapshot schema without modifying it."""
    with path.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t", strict=True)
        if reader.fieldnames != SNAPSHOT_FIELDS:
            raise ValueError("compact snapshot header differs from the required schema")
        rows: list[dict[str, str]] = []
        for raw in reader:
            if None in raw or any(v is None for v in raw.values()):
                raise ValueError("compact snapshot row has the wrong width")
            rows.append({field: raw[field] for field in SNAPSHOT_FIELDS})
    if not rows:
        raise ValueError("compact snapshot contains no records")
    return rows


def check_output(path: Path, inputs: tuple[Path, ...] = ()) -> None:
    """Protect historical source files and all caller-supplied inputs."""
    target = path.resolve()
    if target.is_relative_to(ROOT) or target in {p.resolve() for p in inputs}:
        raise ValueError("output must be separate from historical files and inputs")


def write_table(path: Path, fields: list[str], rows: list[dict[str, str]]) -> None:
    """Atomically replace one explicitly requested TSV after validation."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        mode="w",
        encoding="utf-8",
        newline="",
        dir=path.parent,
        prefix=path.name + ".",
        suffix=".tmp",
        delete=False,
    ) as handle:
        temporary = Path(handle.name)
        try:
            writer = csv.DictWriter(
                handle, fieldnames=fields, dialect="excel-tab", lineterminator="\n"
            )
            writer.writeheader()
            writer.writerows(rows)
        except Exception:
            handle.close()
            temporary.unlink(missing_ok=True)
            raise
    try:
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def select_source(raw: bytes) -> tuple[list[dict[str, str]], int]:
    """Read genuine workbook rows and preserve the original annual selection.

    Offline snapshot builds do not require openpyxl. Source archives are bounded
    before parsing; the reader closes even when a selected record is rejected.
    """
    from openpyxl import load_workbook

    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 64 * 1024 * 1024:
            raise ValueError("official workbook expanded content exceeds the size limit")
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        values = worksheet.iter_rows(values_only=True)
        header = [clean(value) for value in next(values, ())]
        index = {name: position for position, name in enumerate(header)}
        if len(index) != len(header) or not set(RAW_FIELDS).issubset(index):
            raise ValueError(
                "official workbook is missing expected columns or has duplicate headings"
            )
        grouped: dict[str, dict[str, str]] = {}
        raw_data_rows = 0
        for row in values:
            year = clean(row[index["Year"]])
            # Rows 2-4 repeat headings in German, Italian, and French.
            if not year.isdigit():
                continue
            raw_data_rows += 1
            if year != REPORTING_YEAR or clean(row[index["Source type"]]) != "Punktquelle":
                continue
            nace = clean(row[index["NACE code"]])
            if not nace.startswith(NACE_PREFIXES):
                continue
            source_id = clean(row[index["Facility ID"]])
            east_text = clean(row[index["East coordinate (CH1903+)"]])
            north_text = clean(row[index["North coordinate (CH1903+)"]])
            if not source_id.isdigit() or not east_text or not north_text:
                raise ValueError("selected row lacks facility ID or source coordinate")
            east, north = float(east_text), float(north_text)
            if not 2_400_000 <= east <= 2_900_000 or not 1_000_000 <= north <= 1_300_000:
                raise ValueError("selected row has an implausible LV95 coordinate")
            latitude, longitude = lv95_to_wgs84(east, north)
            candidate = {
                "source": "swiss-prtr",
                "source_record_id": source_id,
                "reporting_year": REPORTING_YEAR,
                "facility_name": clean(row[index["Facility"]]),
                "operator": clean(row[index["Owner"]]),
                "country": "Switzerland",
                "lat": f"{latitude:.6f}",
                "lon": f"{longitude:.6f}",
                "source_north_lv95": north_text,
                "source_east_lv95": east_text,
                "nace_code": nace,
                "industrial_sector": clean(row[index["Industrial sector"]]),
                "ordinance_level_2": clean(row[index["Ordinance level 2"]]),
                "ordinance_level_3": clean(row[index["Ordinance level 3"]]),
            }
            previous = grouped.get(source_id)
            if previous is not None and previous != candidate:
                raise ValueError(
                    f"inconsistent identity/classification rows for Facility ID {source_id}"
                )
            grouped[source_id] = candidate
    finally:
        workbook.close()
    selected = sorted(grouped.values(), key=lambda row: int(row["source_record_id"]))
    if not selected:
        raise ValueError("official workbook contains no qualifying records")
    return selected, raw_data_rows


def rows_from_snapshot(
    source_rows: list[dict[str, str]],
    *,
    retrieved: str = RETRIEVED,
) -> list[dict[str, str]]:
    """Preserve reported facility observations without inferring reactor vessels."""
    if date.fromisoformat(retrieved).isoformat() != retrieved:
        raise ValueError("retrieval date must be an exact ISO calendar date")
    if not source_rows:
        raise ValueError("compact snapshot contains no records")
    if any(
        set(row) != set(SNAPSHOT_FIELDS) or any(not isinstance(v, str) for v in row.values())
        for row in source_rows
    ):
        raise ValueError("compact snapshot record differs from the required string schema")
    ids: set[str] = set()
    for row in source_rows:
        for field in (
            "source_record_id",
            "facility_name",
            "operator",
            "lat",
            "lon",
            "nace_code",
            "industrial_sector",
            "source_north_lv95",
            "source_east_lv95",
        ):
            if not row[field].strip():
                raise ValueError(f"compact snapshot is missing {field}")
        if not row["source_record_id"].isdigit():
            raise ValueError("compact snapshot source ID must be numeric")
        north, east = float(row["source_north_lv95"]), float(row["source_east_lv95"])
        if not 1_000_000 <= north <= 1_300_000 or not 2_400_000 <= east <= 2_900_000:
            raise ValueError("compact snapshot LV95 coordinates are out of range")
        expected_lat, expected_lon = lv95_to_wgs84(east, north)
        if row["lat"] != f"{expected_lat:.6f}" or row["lon"] != f"{expected_lon:.6f}":
            raise ValueError("compact snapshot WGS84 point differs from its source LV95 point")
        if row["source_record_id"] in ids:
            raise ValueError("compact snapshot contains a duplicate source ID")
        ids.add(row["source_record_id"])
        if (
            row["source"] != "swiss-prtr"
            or row["reporting_year"] != REPORTING_YEAR
            or row["country"] != "Switzerland"
        ):
            raise ValueError("compact snapshot has an unexpected source, report year or country")
        if not row["nace_code"].startswith(NACE_PREFIXES):
            raise ValueError("compact snapshot NACE code is outside the documented filter")
    rows: list[dict[str, str]] = []
    for src in source_rows:
        process_parts = [src["industrial_sector"]]
        if src["ordinance_level_2"]:
            process_parts.append("PRTRO level 2: " + src["ordinance_level_2"])
        if src["ordinance_level_3"]:
            process_parts.append("PRTRO level 3: " + src["ordinance_level_3"])
        rows.append(
            {
                "stable_id": f"swiss-prtr:{src['source_record_id']}",
                "facility_name": src["facility_name"],
                "country": src["country"],
                "lat": src["lat"],
                "lon": src["lon"],
                "precision": (
                    "SwissPRTR-published CH1903+/LV95 facility point transformed to WGS84 with the "
                    "official swisstopo approximate formula (stated accuracy better than 0.12 arcsec "
                    "longitude and 0.08 arcsec latitude); source position not independently verified"
                ),
                "sector": f"NACE {src['nace_code']} — {src['industrial_sector']}",
                "process_or_activity": " | ".join(process_parts),
                "reactor_type_if_explicit": "",
                "status": (
                    "reported in SwissPRTR for 2024; operating status not asserted; official source "
                    "notes that 2024 release is not yet complete for some facilities"
                ),
                "operator": src["operator"],
                "capacity": "",
                "pollutant_or_product_context": "source industry classification only; pollutant and waste quantities excluded",
                "source_url": LANDING_URL,
                "source_role": "official Swiss pollutant-register facility and activity classification",
                "retrieved": retrieved,
                "license": "Open use; source attribution required (opendata.swiss terms_by)",
                "verification_notes": (
                    f"Selected systematically from 2024 SwissPRTR point sources by NACE prefix in "
                    f"{', '.join(NACE_PREFIXES)}. Published LV95 coordinates were reprojected, not "
                    "geocoded. This record does not establish reactor inventory, design, capacity, "
                    "or current operation."
                ),
            }
        )

    rows.sort(key=lambda row: int(row["stable_id"].split(":", 1)[1]))
    return rows


def refresh_snapshot(
    raw: bytes,
    *,
    snapshot: Path,
    manifest: Path,
    retrieved: str,
) -> tuple[int, int]:
    """Write separate snapshot and raw-source manifest after complete parsing.

    Individual TSV replacements are atomic; the pair is not a multi-file
    transaction. Reporting year derives from the official annual source.
    """
    check_output(snapshot, (manifest,))
    check_output(manifest, (snapshot,))
    selected, raw_rows = select_source(raw)
    rows_from_snapshot(selected, retrieved=retrieved)
    records = [
        {
            "source": "swiss-prtr",
            "url": RAW_URL,
            "retrieved": retrieved,
            "reporting_year": REPORTING_YEAR,
            "bytes": str(len(raw)),
            "sha256": hashlib.sha256(raw).hexdigest(),
            "raw_data_rows": str(raw_rows),
            "selected_facilities": str(len(selected)),
        }
    ]
    write_table(snapshot, SNAPSHOT_FIELDS, selected)
    write_table(manifest, MANIFEST_FIELDS, records)
    return raw_rows, len(selected)


def build(snapshot: Path, output: Path, *, retrieved: str = RETRIEVED) -> int:
    """Rebuild one explicit output from a read-only historical snapshot."""
    check_output(output, (snapshot,))
    rows = rows_from_snapshot(read_snapshot(snapshot), retrieved=retrieved)
    write_table(output, FIELDS, rows)
    return len(rows)


def main(argv: list[str] | None = None) -> int:
    """Build isolated outputs and report invalid source inputs without traceback."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--snapshot", type=Path, default=SNAPSHOT)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--refresh", action="store_true")
    parser.add_argument("--manifest", type=Path)
    parser.add_argument("--date")
    parser.add_argument("--raw-source", type=Path)
    parser.add_argument("--source-url", default=RAW_URL)
    parser.add_argument("--ca-file", type=Path)
    args = parser.parse_args(argv)
    try:
        inputs = tuple(p for p in (args.raw_source, args.ca_file) if p)
        check_output(
            args.output,
            (args.snapshot, *inputs, *((args.manifest,) if args.manifest else ())),
        )
        retrieved = args.date or (date.today().isoformat() if args.refresh else RETRIEVED)
        if args.refresh:
            if args.manifest is None:
                raise ValueError("refresh requires an explicit separate manifest path")
            check_output(args.snapshot, (args.manifest, *inputs))
            check_output(args.manifest, (args.output, args.snapshot, *inputs))
            raw = (
                args.raw_source.read_bytes()
                if args.raw_source
                else download(args.source_url, ca_file=args.ca_file)
            )
            raw_rows, selected_rows = refresh_snapshot(
                raw,
                snapshot=args.snapshot,
                manifest=args.manifest,
                retrieved=retrieved,
            )
            print(f"refreshed snapshot: {raw_rows} raw rows, {selected_rows} selected rows")
        elif args.raw_source or args.manifest:
            raise ValueError("raw source and manifest options require --refresh")
        count = build(args.snapshot, args.output, retrieved=retrieved)
    except (
        OSError,
        UnicodeError,
        csv.Error,
        ValueError,
        ImportError,
        zipfile.BadZipFile,
    ) as exc:
        print(f"FAIL: {exc}")
        return 1
    print(f"wrote {count} records to {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
